import csv
import io
import os
import re
import sqlite3
from typing import List
from urllib.parse import quote_plus

from fastapi import FastAPI, Request, HTTPException
from fastapi.responses import HTMLResponse, JSONResponse, StreamingResponse, Response
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from pydantic import BaseModel

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from openpyxl import Workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont

DB_TABLE = "verses"
COL_ID = "ID"
COL_BOOK = "book"
COL_CHAPTER = "chapter"
COL_VERSE = "verse"
COL_TEXT = "text"
FTS_TABLE = "verses_fts"

BOOK_NAMES = {
    1: "Genesis", 2: "Exodus", 3: "Leviticus", 4: "Numbers", 5: "Deuteronomy",
    6: "Joshua", 7: "Judges", 8: "Ruth", 9: "1 Samuel", 10: "2 Samuel",
    11: "1 Kings", 12: "2 Kings", 13: "1 Chronicles", 14: "2 Chronicles",
    15: "Ezra", 16: "Nehemiah", 17: "Esther", 18: "Job", 19: "Psalms",
    20: "Proverbs", 21: "Ecclesiastes", 22: "Song of Solomon", 23: "Isaiah",
    24: "Jeremiah", 25: "Lamentations", 26: "Ezekiel", 27: "Daniel",
    28: "Hosea", 29: "Joel", 30: "Amos", 31: "Obadiah", 32: "Jonah",
    33: "Micah", 34: "Nahum", 35: "Habakkuk", 36: "Zephaniah", 37: "Haggai",
    38: "Zechariah", 39: "Malachi", 40: "Matthew", 41: "Mark", 42: "Luke",
    43: "John", 44: "Acts", 45: "Romans", 46: "1 Corinthians", 47: "2 Corinthians",
    48: "Galatians", 49: "Ephesians", 50: "Philippians", 51: "Colossians",
    52: "1 Thessalonians", 53: "2 Thessalonians", 54: "1 Timothy", 55: "2 Timothy",
    56: "Titus", 57: "Philemon", 58: "Hebrews", 59: "James", 60: "1 Peter",
    61: "2 Peter", 62: "1 John", 63: "2 John", 64: "3 John", 65: "Jude",
    66: "Revelation"
}

BOOK_NAME_TO_NUM = {v.lower(): k for k, v in BOOK_NAMES.items()}

BOOK_ALIASES = {
    "gen": "genesis", "ge": "genesis", "gn": "genesis",
    "ex": "exodus", "exo": "exodus", "exod": "exodus",
    "lev": "leviticus", "le": "leviticus", "lv": "leviticus",
    "num": "numbers", "nu": "numbers", "nm": "numbers",
    "deut": "deuteronomy", "deu": "deuteronomy", "dt": "deuteronomy",
    "jos": "joshua", "josh": "joshua",
    "jdg": "judges", "jdgs": "judges", "judg": "judges", "jd": "judges",
    "ruth": "ruth", "ru": "ruth",
    "1sam": "1 samuel", "1 sam": "1 samuel", "1sa": "1 samuel",
    "2sam": "2 samuel", "2 sam": "2 samuel", "2sa": "2 samuel",
    "1ki": "1 kings", "1 kgs": "1 kings", "1kgs": "1 kings", "1 kings": "1 kings",
    "2ki": "2 kings", "2 kgs": "2 kings", "2kgs": "2 kings", "2 kings": "2 kings",
    "1chr": "1 chronicles", "1 chr": "1 chronicles", "1ch": "1 chronicles",
    "2chr": "2 chronicles", "2 chr": "2 chronicles", "2ch": "2 chronicles",
    "ezr": "ezra", "ezra": "ezra", "ez": "ezra",
    "neh": "nehemiah", "ne": "nehemiah",
    "est": "esther", "esth": "esther",
    "job": "job",
    "ps": "psalms", "psa": "psalms", "psalm": "psalms", "psalms": "psalms",
    "prov": "proverbs", "pro": "proverbs", "pr": "proverbs", "prv": "proverbs",
    "eccl": "ecclesiastes", "ecc": "ecclesiastes", "ec": "ecclesiastes",
    "song": "song of solomon", "song of songs": "song of solomon",
    "song of solomon": "song of solomon", "sos": "song of solomon",
    "isa": "isaiah", "is": "isaiah",
    "jer": "jeremiah", "je": "jeremiah", "jr": "jeremiah",
    "lam": "lamentations", "la": "lamentations",
    "ezek": "ezekiel", "eze": "ezekiel", "ezk": "ezekiel",
    "dan": "daniel", "da": "daniel", "dn": "daniel",
    "hos": "hosea", "ho": "hosea",
    "joel": "joel", "jl": "joel",
    "amos": "amos", "am": "amos",
    "obad": "obadiah", "ob": "obadiah", "oba": "obadiah",
    "jon": "jonah", "jnh": "jonah",
    "mic": "micah", "mc": "micah",
    "nah": "nahum", "na": "nahum",
    "hab": "habakkuk", "hb": "habakkuk",
    "zeph": "zephaniah", "zep": "zephaniah", "zp": "zephaniah",
    "hag": "haggai", "hg": "haggai",
    "zech": "zechariah", "zec": "zechariah", "zc": "zechariah",
    "mal": "malachi", "ml": "malachi",
    "mat": "matthew", "matt": "matthew", "mt": "matthew",
    "mrk": "mark", "mar": "mark", "mk": "mark", "mr": "mark",
    "luk": "luke", "lk": "luke", "lu": "luke",
    "jhn": "john", "jn": "john", "joh": "john",
    "act": "acts", "acts": "acts", "ac": "acts",
    "rom": "romans", "rm": "romans", "ro": "romans",
    "1cor": "1 corinthians", "1 cor": "1 corinthians", "1co": "1 corinthians",
    "2cor": "2 corinthians", "2 cor": "2 corinthians", "2co": "2 corinthians",
    "gal": "galatians", "ga": "galatians",
    "eph": "ephesians", "ep": "ephesians",
    "php": "philippians", "phil": "philippians", "phl": "philippians",
    "col": "colossians", "co": "colossians",
    "1th": "1 thessalonians", "1 th": "1 thessalonians", "1thes": "1 thessalonians", "1thess": "1 thessalonians",  "1 thess": "1 thessalonians", "1 thes": "1 thessalonians",
    "2th": "2 thessalonians", "2 th": "2 thessalonians", "2thes": "2 thessalonians", "2thess": "2 thessalonians",  "2 thess": "2 thessalonians", "2 thes": "2 thessalonians",
    "1tim": "1 timothy", "1 tim": "1 timothy", "1ti": "1 timothy",
    "2tim": "2 timothy", "2 tim": "2 timothy", "2ti": "2 timothy",
    "tit": "titus", "ti": "titus",
    "phm": "philemon", "philem": "philemon", "pm": "philemon",
    "heb": "hebrews", "he": "hebrews",
    "jas": "james", "jam": "james", "jm": "james", "js": "james",
    "1pet": "1 peter", "1 pet": "1 peter", "1pe": "1 peter", "1pt": "1 peter", "1 pt": "1 peter", "1peter": "1 peter",
    "2pet": "2 peter", "2 pet": "2 peter", "2pe": "2 peter", "2pt": "2 peter", "2 pt": "2 peter", "2peter": "2 peter",
    "1jn": "1 john", "1 jn": "1 john", "1jhn": "1 john", "1jo": "1 john", "1j": "1 john",
    "2jn": "2 john", "2 jn": "2 john", "2jhn": "2 john", "2jo": "2 john", "2j": "2 john",
    "3jn": "3 john", "3 jn": "3 john", "3jhn": "3 john", "3jo": "3 john", "3j": "3 john",
    "jud": "jude", "jude": "jude",
    "rev": "revelation", "re": "revelation", "rv": "revelation"
}

BOOK_GROUPS = {
    "All Books": list(range(1, 67)),
    "Law": [1, 2, 3, 4, 5],
    "Prophets": [23, 24, 25, 26, 27, 28, 29, 30, 31, 32, 33, 34, 35, 36, 37, 38, 39],
    "Psalms": [19],
    "Gospels": [40, 41, 42, 43],
    "Epistles": [45, 46, 47, 48, 49, 50, 51, 52, 53, 54, 55, 56, 57, 58, 59, 60, 61, 62, 63, 64, 65],
}

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DB_PATH = os.path.join(BASE_DIR, "kjv.sqlite")

app = FastAPI(title="KJV Bible Search Web")
app.mount("/static", StaticFiles(directory=os.path.join(BASE_DIR, "static")), name="static")
templates = Jinja2Templates(directory=os.path.join(BASE_DIR, "templates"))


class SearchRequest(BaseModel):
    intent: str = "search"
    query: str
    book_group: str = "All Books"
    search_type: str = "Phrase"
    whole_word: bool = False
    case_sensitive: bool = False
    reference_style: str = "Wikilink"
    granularity: str = "Verse"
    reference_position: str = "Ref First"
    highlight: bool = True
    suppress_paragraph: bool = False
    lookup_extracted: bool = True
    context_window: int = 0
    show_red_letters: bool = True


def get_conn():
    if not os.path.exists(DB_PATH):
        raise FileNotFoundError(f"Database not found: {DB_PATH}")
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn


def normalize_book_token(token):
    token = token.strip().lower()
    token = re.sub(r"\s+", " ", token)
    token = token.replace(".", "")
    return token


def resolve_book_name(book_part):
    token = normalize_book_token(book_part)
    if token in BOOK_NAME_TO_NUM:
        return token
    if token in BOOK_ALIASES:
        return BOOK_ALIASES[token]

    matches = [name for name in BOOK_NAME_TO_NUM.keys() if name.startswith(token)]
    if len(matches) == 1:
        return matches[0]

    compact = token.replace(" ", "")
    matches = [name for name in BOOK_NAME_TO_NUM.keys() if name.replace(" ", "").startswith(compact)]
    if len(matches) == 1:
        return matches[0]

    return None


def title_case_book_name(name):
    if name == "psalms":
        return "Psalms"
    return " ".join(part.capitalize() for part in name.split())


def parse_reference(ref_text):
    ref_text = ref_text.strip()

    resolved_book = resolve_book_name(ref_text)
    if resolved_book:
        return {
            "book_name": title_case_book_name(resolved_book),
            "book_num": BOOK_NAME_TO_NUM[resolved_book],
            "chapter": None,
            "verse_start": None,
            "verse_end": None
        }

    pattern = r"^\s*(.+?)\s+(\d+)(?:[:.](\d+)(?:-(\d+))?)?\s*$"
    m = re.match(pattern, ref_text, re.IGNORECASE)
    if not m:
        return None

    book_part = m.group(1).strip()
    chapter = int(m.group(2))
    verse_start = int(m.group(3)) if m.group(3) else None
    verse_end = int(m.group(4)) if m.group(4) else None

    resolved_book = resolve_book_name(book_part)
    if not resolved_book:
        return None

    return {
        "book_name": title_case_book_name(resolved_book),
        "book_num": BOOK_NAME_TO_NUM[resolved_book],
        "chapter": chapter,
        "verse_start": verse_start,
        "verse_end": verse_end
    }


def split_reference_tokens(text):
    text = text.replace(";", " ")
    text = text.replace(",", " ")
    text = re.sub(r"\s+", " ", text.strip())
    if not text:
        return []
    return text.split()


def try_parse_book_from_tokens(tokens, start_idx):
    max_take = min(3, len(tokens) - start_idx)
    for take in range(max_take, 0, -1):
        candidate = " ".join(tokens[start_idx:start_idx + take])
        resolved = resolve_book_name(candidate)
        if resolved:
            return resolved, start_idx + take
    return None, start_idx


def try_parse_compact_book_token(token):
    raw_token = token.strip().lower()
    raw_token = re.sub(r"\s+", "", raw_token)

    candidates = sorted(
        set(list(BOOK_ALIASES.keys()) + list(BOOK_NAME_TO_NUM.keys())),
        key=len,
        reverse=True
    )

    for candidate in candidates:
        compact_candidate = candidate.lower().replace(" ", "").replace(".", "")
        if raw_token.startswith(compact_candidate):
            remainder = raw_token[len(compact_candidate):]
            if remainder and re.match(r"^\d+(?:[:.]\d+(?:-\d+)?)?$", remainder):
                resolved = resolve_book_name(candidate)
                if resolved:
                    return resolved, remainder

    return None, None


def parse_chapter_verse_token(token):
    token = token.strip()

    m = re.match(r"^(\d+)[:.](\d+)(?:-(\d+))?$", token)
    if m:
        return {
            "kind": "chapter_verse",
            "chapter": int(m.group(1)),
            "verse_start": int(m.group(2)),
            "verse_end": int(m.group(3)) if m.group(3) else None
        }

    m = re.match(r"^(\d+)-(\d+)$", token)
    if m:
        return {
            "kind": "verse_only",
            "verse_start": int(m.group(1)),
            "verse_end": int(m.group(2))
        }

    m = re.match(r"^(\d+)$", token)
    if m:
        return {
            "kind": "single_number",
            "value": int(m.group(1))
        }

    return None


def parse_reference_list(ref_text):
    ref_text = ref_text.strip()
    if not ref_text:
        return [], None

    single = parse_reference(ref_text)
    if single:
        return [single], None

    tokens = split_reference_tokens(ref_text)
    if not tokens:
        return [], None

    parsed_refs = []
    current_book = None
    current_chapter = None
    last_ref_type = None

    i = 0
    while i < len(tokens):
        token = tokens[i]

        resolved_book, next_i = try_parse_book_from_tokens(tokens, i)
        if resolved_book:
            current_book = resolved_book
            current_chapter = None
            last_ref_type = "book"
            i = next_i

            if i >= len(tokens):
                parsed_refs.append({
                    "book_name": title_case_book_name(current_book),
                    "book_num": BOOK_NAME_TO_NUM[current_book],
                    "chapter": None,
                    "verse_start": None,
                    "verse_end": None
                })
                break

            next_token = tokens[i]
            token_info = parse_chapter_verse_token(next_token)

            if token_info is None:
                parsed_refs.append({
                    "book_name": title_case_book_name(current_book),
                    "book_num": BOOK_NAME_TO_NUM[current_book],
                    "chapter": None,
                    "verse_start": None,
                    "verse_end": None
                })
                continue

            token = next_token
        else:
            compact_book, remainder = try_parse_compact_book_token(token)
            if compact_book:
                current_book = compact_book
                current_chapter = None
                last_ref_type = "book"

                token_info = parse_chapter_verse_token(remainder)
                if token_info is None:
                    return [], token

                if token_info["kind"] == "chapter_verse":
                    current_chapter = token_info["chapter"]
                    parsed_refs.append({
                        "book_name": title_case_book_name(current_book),
                        "book_num": BOOK_NAME_TO_NUM[current_book],
                        "chapter": current_chapter,
                        "verse_start": token_info["verse_start"],
                        "verse_end": token_info["verse_end"]
                    })
                    last_ref_type = "verse"
                    i += 1
                    continue

                if token_info["kind"] == "single_number":
                    current_chapter = token_info["value"]
                    parsed_refs.append({
                        "book_name": title_case_book_name(current_book),
                        "book_num": BOOK_NAME_TO_NUM[current_book],
                        "chapter": current_chapter,
                        "verse_start": None,
                        "verse_end": None
                    })
                    last_ref_type = "chapter"
                    i += 1
                    continue

                return [], token

        if current_book is None:
            return [], token

        token_info = parse_chapter_verse_token(token)
        if token_info is None:
            return [], token

        if token_info["kind"] == "chapter_verse":
            current_chapter = token_info["chapter"]
            parsed_refs.append({
                "book_name": title_case_book_name(current_book),
                "book_num": BOOK_NAME_TO_NUM[current_book],
                "chapter": current_chapter,
                "verse_start": token_info["verse_start"],
                "verse_end": token_info["verse_end"]
            })
            last_ref_type = "verse"
            i += 1
            continue

        if token_info["kind"] == "verse_only":
            if current_chapter is None:
                return [], token

            parsed_refs.append({
                "book_name": title_case_book_name(current_book),
                "book_num": BOOK_NAME_TO_NUM[current_book],
                "chapter": current_chapter,
                "verse_start": token_info["verse_start"],
                "verse_end": token_info["verse_end"]
            })
            last_ref_type = "verse"
            i += 1
            continue

        if token_info["kind"] == "single_number":
            value = token_info["value"]

            if last_ref_type == "verse":
                if current_chapter is None:
                    return [], token

                parsed_refs.append({
                    "book_name": title_case_book_name(current_book),
                    "book_num": BOOK_NAME_TO_NUM[current_book],
                    "chapter": current_chapter,
                    "verse_start": value,
                    "verse_end": None
                })
                last_ref_type = "verse"
            else:
                current_chapter = value
                parsed_refs.append({
                    "book_name": title_case_book_name(current_book),
                    "book_num": BOOK_NAME_TO_NUM[current_book],
                    "chapter": current_chapter,
                    "verse_start": None,
                    "verse_end": None
                })
                last_ref_type = "chapter"

            i += 1
            continue

    return parsed_refs, None


def get_all_book_tokens():
    tokens = set(BOOK_NAME_TO_NUM.keys()) | set(BOOK_ALIASES.keys())
    return sorted(tokens, key=len, reverse=True)


def build_book_regex():
    parts = []
    for token in get_all_book_tokens():
        escaped = re.escape(token)
        escaped = escaped.replace(r"\ ", r"\s*")
        parts.append(escaped)
    return "(?:" + "|".join(parts) + ")"


def extract_wikilink_candidates(text):
    return re.findall(r"\[\[([^\]]+)\]\]", text)


def extract_plain_reference_candidates(text):
    book_re = build_book_regex()

    pattern = re.compile(
        rf"""
        (
            {book_re}
            \s*
            \d+
            (?:
                [:.]\d+(?:-\d+)?
            )?
            (?:
                \s*[,;]\s*
                (?:
                    \d+(?:[:.]\d+(?:-\d+)?)?
                )
            )*
        )
        """,
        re.IGNORECASE | re.VERBOSE
    )

    return [m.group(1).strip() for m in pattern.finditer(text)]


def extract_references_from_text(text):
    candidates = []
    candidates.extend(extract_wikilink_candidates(text))
    candidates.extend(extract_plain_reference_candidates(text))

    parsed_refs = []
    seen = set()

    for candidate in candidates:
        parsed_list, bad_token = parse_reference_list(candidate)
        if parsed_list and not bad_token:
            for parsed in parsed_list:
                key = (
                    parsed["book_num"],
                    parsed["chapter"],
                    parsed["verse_start"],
                    parsed["verse_end"]
                )
                if key not in seen:
                    seen.add(key)
                    parsed_refs.append(parsed)

    return parsed_refs


def resolve_reference_intent(query):
    query = query.strip()

    single = parse_reference(query)
    if single:
        return {
            "mode": "reference",
            "references": [single],
            "recommended_view": "reading"
        }

    parsed_refs, bad_token = parse_reference_list(query)
    if parsed_refs and not bad_token:
        return {
            "mode": "reference",
            "references": parsed_refs,
            "recommended_view": "reading" if len(parsed_refs) == 1 else "results"
        }

    extracted = extract_references_from_text(query)
    if extracted:
        return {
            "mode": "extract",
            "references": extracted,
            "recommended_view": "reading" if len(extracted) == 1 else "results"
        }

    return {
        "mode": "reference",
        "references": [],
        "recommended_view": "results"
    }


def wildcard_to_regex(term):
    escaped = re.escape(term)
    if "*" in term:
        return escaped.replace(r"\*", r"\w*")
    return escaped


def clean_text(text, suppress_paragraph):
    if suppress_paragraph:
        text = text.replace("¶", "")
        text = re.sub(r"\s{2,}", " ", text).strip()
    return text


def split_paragraph_parts(text, suppress_paragraph=False):
    parts = []
    for raw_part in str(text).split("¶"):
        clean_part = clean_text(raw_part, suppress_paragraph)
        if clean_part:
            parts.append(clean_part)
    return parts


def compile_search_regexes(search_text, search_mode, whole_word, case_sensitive):
    if not search_text:
        return []

    flags = 0 if case_sensitive else re.IGNORECASE
    regexes = []

    def word_pattern(word):
        pattern = wildcard_to_regex(word)
        return r"\b" + pattern + r"\b"

    if search_mode == "Phrase":
        words = search_text.split()
        if not words:
            return []

        pattern_text = r"\s+".join(wildcard_to_regex(word) for word in words)
        regexes.append(re.compile(r"\b" + pattern_text + r"\b", flags))
    else:
        words = search_text.split()
        for word in words:
            regexes.append(re.compile(word_pattern(word), flags))

    return regexes


def get_selected_book_numbers(book_group):
    return BOOK_GROUPS.get(book_group, list(range(1, 67)))


def has_fts_table(cur):
    cur.execute("""
        SELECT name
        FROM sqlite_master
        WHERE type='table' AND name=?
    """, (FTS_TABLE,))
    return cur.fetchone() is not None


def build_fts_query(search_text, search_mode):
    if search_mode == "Phrase":
        escaped = search_text.replace('"', '""')
        return f'"{escaped}"'

    words = search_text.split()
    cleaned_words = []
    for word in words:
        cleaned = word.replace('"', "").strip().replace("’", "'")
        cleaned = re.sub(r"[^\w*']", "", cleaned)
        if cleaned:
            cleaned_words.append(cleaned)

    if not cleaned_words:
        return ""

    if search_mode == "All Words":
        return " AND ".join(cleaned_words)
    return " OR ".join(cleaned_words)


def can_use_fts(cur, search_text, case_sensitive):
    if case_sensitive:
        return False
    if "*" in search_text:
        return False
    if not search_text.strip():
        return False
    if "'" in search_text or "’" in search_text:
        return False
    if re.search(r"[:.\-]", search_text):
        return False
    return has_fts_table(cur)


def get_fts_text_rows(cur, search_text, selected_books, search_mode):
    placeholders = ",".join("?" for _ in selected_books)
    fts_query = build_fts_query(search_text, search_mode)

    query = f"""
        SELECT v.{COL_ID}, v.{COL_BOOK}, v.{COL_CHAPTER}, v.{COL_VERSE}, v.{COL_TEXT}
        FROM {FTS_TABLE} f
        JOIN {DB_TABLE} v ON v.{COL_ID} = f.rowid
        WHERE f.{COL_TEXT} MATCH ?
          AND v.{COL_BOOK} IN ({placeholders})
        ORDER BY v.{COL_ID}
    """

    params = [fts_query] + list(selected_books)
    cur.execute(query, params)
    return cur.fetchall()


def get_sqlite_text_rows(cur, search_text, selected_books, search_mode, case_sensitive):
    placeholders = ",".join("?" for _ in selected_books)

    base_query = f"""
        SELECT {COL_ID}, {COL_BOOK}, {COL_CHAPTER}, {COL_VERSE}, {COL_TEXT}
        FROM {DB_TABLE}
        WHERE {COL_BOOK} IN ({placeholders})
    """

    params = list(selected_books)

    if "*" in search_text:
        base_query += f" ORDER BY {COL_ID}"
        cur.execute(base_query, params)
        return cur.fetchall()

    literal_apostrophe_query = "'" in search_text or "’" in search_text

    if search_mode == "Phrase":
        if literal_apostrophe_query:
            base_query += f" AND ({COL_TEXT} LIKE ? ESCAPE '\\')"
            params.append(f"%{search_text}%")
        elif case_sensitive:
            base_query += f" AND instr({COL_TEXT}, ?) > 0"
            params.append(search_text)
        else:
            base_query += f" AND instr(lower({COL_TEXT}), lower(?)) > 0"
            params.append(search_text)
    else:
        words = search_text.split()
        if not words:
            return []

        joiner = " AND " if search_mode == "All Words" else " OR "
        clauses = []

        for word in words:
            if literal_apostrophe_query:
                clauses.append(f"({COL_TEXT} LIKE ? ESCAPE '\\')")
                params.append(f"%{word}%")
            elif case_sensitive:
                clauses.append(f"instr({COL_TEXT}, ?) > 0")
                params.append(word)
            else:
                clauses.append(f"instr(lower({COL_TEXT}), lower(?)) > 0")
                params.append(word)

        base_query += " AND (" + joiner.join(clauses) + ")"

    base_query += f" ORDER BY {COL_ID}"
    cur.execute(base_query, params)
    return cur.fetchall()


def filter_rows_python(rows, regexes, search_mode):
    if not regexes:
        return [], 0

    filtered = []
    total_matches = 0

    if search_mode == "Phrase":
        regex = regexes[0]
        for row in rows:
            matches = list(regex.finditer(row["cleaned_text"]))
            if matches:
                filtered.append(row)
                total_matches += len(matches)
    else:
        for row in rows:
            verse_text = row["cleaned_text"]
            match_lists = [list(rx.finditer(verse_text)) for rx in regexes]
            bools = [bool(m) for m in match_lists]

            if search_mode == "All Words":
                if all(bools):
                    filtered.append(row)
                    total_matches += sum(len(m) for m in match_lists)
            else:
                if any(bools):
                    filtered.append(row)
                    total_matches += sum(len(m) for m in match_lists)

    return filtered, total_matches


def get_reference_rows(cur, parsed):
    book_num = parsed["book_num"]
    chapter = parsed["chapter"]
    verse_start = parsed["verse_start"]
    verse_end = parsed["verse_end"]

    if chapter is None:
        query = f"""
            SELECT {COL_ID}, {COL_BOOK}, {COL_CHAPTER}, {COL_VERSE}, {COL_TEXT}
            FROM {DB_TABLE}
            WHERE {COL_BOOK} = ?
            ORDER BY {COL_ID}
        """
        params = (book_num,)
    elif verse_start is None:
        query = f"""
            SELECT {COL_ID}, {COL_BOOK}, {COL_CHAPTER}, {COL_VERSE}, {COL_TEXT}
            FROM {DB_TABLE}
            WHERE {COL_BOOK} = ?
              AND {COL_CHAPTER} = ?
            ORDER BY {COL_ID}
        """
        params = (book_num, chapter)
    elif verse_end is None:
        query = f"""
            SELECT {COL_ID}, {COL_BOOK}, {COL_CHAPTER}, {COL_VERSE}, {COL_TEXT}
            FROM {DB_TABLE}
            WHERE {COL_BOOK} = ?
              AND {COL_CHAPTER} = ?
              AND {COL_VERSE} = ?
        """
        params = (book_num, chapter, verse_start)
    else:
        query = f"""
            SELECT {COL_ID}, {COL_BOOK}, {COL_CHAPTER}, {COL_VERSE}, {COL_TEXT}
            FROM {DB_TABLE}
            WHERE {COL_BOOK} = ?
              AND {COL_CHAPTER} = ?
              AND {COL_VERSE} BETWEEN ? AND ?
            ORDER BY {COL_ID}
        """
        params = (book_num, chapter, verse_start, verse_end)

    cur.execute(query, params)
    return cur.fetchall()


def fetch_context_rows(cur, base_rows, context_window):
    if context_window <= 0 or not base_rows:
        return base_rows

    ids = set()

    for row in base_rows:
        if isinstance(row, sqlite3.Row):
            row_id = int(row[COL_ID])
        else:
            row_id = int(row["id"])
        for i in range(row_id - context_window, row_id + context_window + 1):
            if i > 0:
                ids.add(i)

    if not ids:
        return base_rows

    placeholders = ",".join("?" for _ in ids)
    query = f"""
        SELECT {COL_ID}, {COL_BOOK}, {COL_CHAPTER}, {COL_VERSE}, {COL_TEXT}
        FROM {DB_TABLE}
        WHERE {COL_ID} IN ({placeholders})
        ORDER BY {COL_ID}
    """
    cur.execute(query, list(sorted(ids)))
    return cur.fetchall()


def get_book_name(book_num):
    return BOOK_NAMES.get(int(book_num), str(book_num))


def get_single_reference(book_num, chapter, verse, reference_style):
    book_name = get_book_name(book_num)
    if reference_style == "Wikilink":
        return f"[[{book_name} {chapter}.{verse}]]"
    return f"{book_name} {chapter}:{verse}"


def get_range_reference(book_num, chapter, verse_start, verse_end, reference_style):
    book_name = get_book_name(book_num)
    if verse_start == verse_end:
        return get_single_reference(book_num, chapter, verse_start, reference_style)
    if reference_style == "Wikilink":
        return f"[[{book_name} {chapter}.{verse_start}-{verse_end}]]"
    return f"{book_name} {chapter}:{verse_start}-{verse_end}"


def parse_red_letter_segments(text, show_red_letters=True):
    if not show_red_letters:
        clean = text.replace("‹", "").replace("›", "")
        return [{"text": clean, "red": False}]

    segments = []
    pos = 0
    pattern = re.compile(r"‹(.*?)›")

    for m in pattern.finditer(text):
        if m.start() > pos:
            segments.append({"text": text[pos:m.start()], "red": False})
        segments.append({"text": m.group(1), "red": True})
        pos = m.end()

    if pos < len(text):
        segments.append({"text": text[pos:], "red": False})

    if not segments:
        return [{"text": text.replace("‹", "").replace("›", ""), "red": False}]

    cleaned = []
    for seg in segments:
        cleaned.append({
            "text": seg["text"].replace("‹", "").replace("›", ""),
            "red": seg["red"]
        })
    return cleaned


def get_highlight_spans(text, regexes, do_highlight):
    if not do_highlight:
        return []

    spans = []
    for rx in regexes:
        for m in rx.finditer(text):
            if m.start() != m.end():
                spans.append((m.start(), m.end()))

    if not spans:
        return []

    spans.sort()
    merged = [list(spans[0])]
    for start, end in spans[1:]:
        if start <= merged[-1][1]:
            merged[-1][1] = max(merged[-1][1], end)
        else:
            merged.append([start, end])

    return [(s, e) for s, e in merged]


def normalize_row(row, suppress_paragraph):
    return {
        "id": int(row[COL_ID]),
        "book": int(row[COL_BOOK]),
        "chapter": int(row[COL_CHAPTER]),
        "verse": int(row[COL_VERSE]),
        "text": row[COL_TEXT],
        "cleaned_text": clean_text(row[COL_TEXT], suppress_paragraph),
        "suppress_paragraph": bool(suppress_paragraph)
    }


def build_display_items(rows, granularity, reference_style, reference_position, regexes, do_highlight, show_red_letters):
    if granularity == "Verse":
        return build_verse_items(rows, reference_style, reference_position, regexes, do_highlight, show_red_letters)
    return build_passage_items(rows, reference_style, reference_position, regexes, do_highlight, show_red_letters)


def build_verse_items(rows, reference_style, reference_position, regexes, do_highlight, show_red_letters):
    items = []
    for row in rows:
        reference = get_single_reference(row["book"], row["chapter"], row["verse"], reference_style)
        items.append({
            "reference": reference,
            "verse_text": row["cleaned_text"].replace("‹", "").replace("›", ""),
            "book": row["book"],
            "chapter": row["chapter"],
            "verse": row["verse"],
            "reference_position": reference_position,
            "highlight_spans": get_highlight_spans(row["cleaned_text"].replace("‹", "").replace("›", ""), regexes, do_highlight),
            "segments": parse_red_letter_segments(row["cleaned_text"], show_red_letters)
        })
    return items


def build_passage_items(rows, reference_style, reference_position, regexes, do_highlight, show_red_letters):
    if not rows:
        return []

    groups = []
    current_group = [rows[0]]

    for row in rows[1:]:
        prev = current_group[-1]
        if (
            row["book"] == prev["book"]
            and row["chapter"] == prev["chapter"]
            and row["verse"] == prev["verse"] + 1
        ):
            current_group.append(row)
        else:
            groups.append(current_group)
            current_group = [row]
    groups.append(current_group)

    items = []
    for group in groups:
        book = group[0]["book"]
        chapter = group[0]["chapter"]
        verse_start = group[0]["verse"]
        verse_end = group[-1]["verse"]

        paragraph_segments = []
        for row in group:
            for part in split_paragraph_parts(row["text"], row.get("suppress_paragraph", False)):
                paragraph_segments.append(part)

        passage_text_raw = "\n\n".join(paragraph_segments).strip()
        passage_text_clean = passage_text_raw.replace("‹", "").replace("›", "")
        reference = get_range_reference(book, chapter, verse_start, verse_end, reference_style)

        items.append({
            "reference": reference,
            "verse_text": passage_text_clean,
            "book": book,
            "chapter": chapter,
            "verse": verse_start,
            "reference_position": reference_position,
            "highlight_spans": get_highlight_spans(passage_text_clean, regexes, do_highlight),
            "segments": parse_red_letter_segments(passage_text_raw, show_red_letters)
        })
    return items


def clean_reading_text(text):
    text = text.strip()
    text = text.replace("‹", "").replace("›", "")
    text = re.sub(r'^[\s"\']+', '', text)
    return text.strip()


def split_paragraphs_from_rows(rows, show_red_letters):
    paragraphs = []
    current = []

    for row in rows:
        raw_text = row.get("text", row.get("cleaned_text", ""))
        parts = str(raw_text).split("¶")

        for idx, part in enumerate(parts):
            if idx > 0 and current and "¶" in str(row.get("text", row.get("cleaned_text", ""))):
                paragraphs.append(current)
                current = []

            part = clean_text(part, row.get("suppress_paragraph", False)).strip()
            if not part:
                continue

            segments = parse_red_letter_segments(part, show_red_letters)
            clean_plain = clean_reading_text("".join(seg["text"] for seg in segments))

            if clean_plain:
                current.append({
                    "verse": row["verse"],
                    "text": clean_plain,
                    "segments": segments
                })

            if idx < len(parts) - 1:
                if current:
                    paragraphs.append(current)
                    current = []

    if current:
        paragraphs.append(current)

    return paragraphs


def chapter_exists(cur, book, chapter):
    cur.execute(
        f"""
        SELECT 1
        FROM {DB_TABLE}
        WHERE {COL_BOOK}=? AND {COL_CHAPTER}=?
        LIMIT 1
        """,
        (book, chapter)
    )
    return cur.fetchone() is not None


def get_adjacent_chapter_refs(cur, book, chapter):
    prev_ref = None
    next_ref = None

    prev_book = book
    prev_chapter = chapter - 1
    if prev_chapter < 1:
        prev_book = book - 1
        while prev_book >= 1:
            cur.execute(
                f"SELECT MAX({COL_CHAPTER}) AS max_ch FROM {DB_TABLE} WHERE {COL_BOOK}=?",
                (prev_book,)
            )
            row = cur.fetchone()
            if row and row["max_ch"]:
                prev_chapter = int(row["max_ch"])
                prev_ref = f"{get_book_name(prev_book)} {prev_chapter}"
                break
            prev_book -= 1
    else:
        if chapter_exists(cur, book, prev_chapter):
            prev_ref = f"{get_book_name(book)} {prev_chapter}"

    next_book = book
    next_chapter = chapter + 1
    if chapter_exists(cur, book, next_chapter):
        next_ref = f"{get_book_name(book)} {next_chapter}"
    else:
        next_book = book + 1
        while next_book <= 66:
            if chapter_exists(cur, next_book, 1):
                next_ref = f"{get_book_name(next_book)} 1"
                break
            next_book += 1

    return {"previous": prev_ref, "next": next_ref}


def build_reading_items(rows, reference_style, show_red_letters, cur):
    if not rows:
        return []

    groups = []
    current_group = [rows[0]]

    for row in rows[1:]:
        prev = current_group[-1]
        if (
            row["book"] == prev["book"]
            and row["chapter"] == prev["chapter"]
            and row["verse"] == prev["verse"] + 1
        ):
            current_group.append(row)
        else:
            groups.append(current_group)
            current_group = [row]
    groups.append(current_group)

    reading_items = []

    for group in groups:
        book = group[0]["book"]
        chapter = group[0]["chapter"]
        verse_start = group[0]["verse"]
        verse_end = group[-1]["verse"]

        reference = get_range_reference(book, chapter, verse_start, verse_end, reference_style)
        paragraphs = split_paragraphs_from_rows(group, show_red_letters)

        reading_items.append({
            "reference": reference,
            "book": book,
            "chapter": chapter,
            "nav": get_adjacent_chapter_refs(cur, book, chapter),
            "paragraphs": paragraphs
        })

    return reading_items


def markdown_escape_minimal(text):
    return str(text).replace("\\", "\\\\").replace("*", r"\*").replace("_", r"\_")


def is_highlight_overlap(start, end, highlight_spans):
    for span_start, span_end in highlight_spans:
        if start < span_end and end > span_start:
            return True
    return False


def build_text_segments(text, spans):
    if not spans:
        return [(text, False)]

    segments = []
    last = 0
    for start, end in spans:
        start = max(0, min(start, len(text)))
        end = max(start, min(end, len(text)))
        if start > last:
            segments.append((text[last:start], False))
        if end > start:
            segments.append((text[start:end], True))
        last = max(last, end)

    if last < len(text):
        segments.append((text[last:], False))

    return segments


def build_relative_spans_for_slice(text, offset, highlight_spans):
    if not highlight_spans:
        return []
    local = []
    for span_start, span_end in highlight_spans:
        local_start = span_start - offset
        local_end = span_end - offset
        if local_end <= 0 or local_start >= len(text):
            continue
        local.append((max(0, local_start), min(len(text), local_end)))
    return local


def build_bracketed_segments(text, highlight_spans=None):
    highlight_spans = highlight_spans or []
    segments = []
    pattern = re.compile(r"\[([^\]]+)\]")
    match_found = False
    last_index = 0

    for match in pattern.finditer(text):
        match_found = True
        prefix = text[last_index:match.start()]
        if prefix:
            relative = build_relative_spans_for_slice(prefix, last_index, highlight_spans)
            for seg_text, is_match in build_text_segments(prefix, relative):
                segments.append((seg_text, is_match, False))

        inner = match.group(1)
        if inner:
            inner_raw_start = match.start() + 1
            inner_raw_end = inner_raw_start + len(inner)
            local_spans = []
            for span_start, span_end in highlight_spans:
                if span_end <= match.start() or span_start >= match.end():
                    continue
                offset = inner_raw_start if span_start >= inner_raw_start else match.start()
                local_start = max(0, span_start - offset)
                local_end = max(local_start, min(len(inner), span_end - offset))
                if local_end > local_start:
                    local_spans.append((local_start, local_end))

            if local_spans:
                for seg_text, is_match in build_text_segments(inner, local_spans):
                    segments.append((seg_text, is_match, True))
            else:
                segments.append((inner, False, True))

        last_index = match.end()

    if last_index < len(text):
        suffix = text[last_index:]
        if suffix:
            relative = build_relative_spans_for_slice(suffix, last_index, highlight_spans)
            for seg_text, is_match in build_text_segments(suffix, relative):
                segments.append((seg_text, is_match, False))

    if not match_found:
        relative = build_relative_spans_for_slice(text, 0, highlight_spans)
        return [(seg_text, is_match, False) for seg_text, is_match in build_text_segments(text, relative)]

    return segments


def get_export_verse_prefix(item):
    reference = item.get("reference") or ""
    if reference:
        return ""

    verse = item.get("verse")
    if verse is None or verse == "":
        return ""
    return f"{verse}. "


def build_plain_text(item):
    reference = item.get("reference") or ""
    verse_text = item.get("verse_text") or ""
    verse_prefix = get_export_verse_prefix(item)

    if item.get("reference_position") == "Ref Last":
        if not reference:
            return f"{verse_prefix}{verse_text}".strip()
        if verse_prefix:
            return f"{verse_prefix}{verse_text} {reference}".strip()
        return f"{verse_text} {reference}".strip()

    if not reference:
        return f"{verse_prefix}{verse_text}".strip()
    if verse_prefix:
        return f"{reference} {verse_prefix}{verse_text}".strip()
    return f"{reference} {verse_text}".strip()


def build_markdown_text(item, base_url: str | None = None):
    segments = build_bracketed_segments(item["verse_text"], item.get("highlight_spans", []))
    reference = item.get("reference") or ""
    reference_position = item.get("reference_position") or "Ref First"
    verse_prefix = get_export_verse_prefix(item) if not reference else ""

    parts = []
    for seg_text, is_match, is_bracketed in segments:
        seg_text = markdown_escape_minimal(seg_text)
        if is_match:
            parts.append(f"***{seg_text}***")
        elif is_bracketed:
            parts.append(f"*{seg_text}*")
        else:
            parts.append(seg_text)

    verse_text = "".join(parts)
    if verse_prefix:
        verse_text = f"{verse_prefix}{verse_text}"

    markdown_reference = build_markdown_reading_link(reference, base_url=base_url)
    web_link = build_markdown_web_link(reference, base_url=base_url)
    if not markdown_reference:
        return f"{verse_text} {web_link}".strip()

    if reference_position == "Ref Last":
        return f"{verse_text} {markdown_reference} {web_link}".strip()
    return f"{markdown_reference} {verse_text} {web_link}".strip()


def add_docx_hyperlink(paragraph, text, url):
    if not text or not url:
        paragraph.add_run(text or "")
        return

    part = paragraph.part
    rel_id = part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)

    run = OxmlElement("w:r")
    run_props = OxmlElement("w:rPr")
    style = OxmlElement("w:rStyle")
    style.set(qn("w:val"), "Hyperlink")
    run_props.append(style)
    run.append(run_props)

    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def write_docx_runs(paragraph, item, base_url: str | None = None):
    segments = build_bracketed_segments(item["verse_text"], item.get("highlight_spans", []))
    verse_prefix = get_export_verse_prefix(item)
    reference = item.get("reference") or ""
    reference_position = item.get("reference_position") or "Ref First"

    if reference_position == "Ref Last":
        if verse_prefix:
            paragraph.add_run(verse_prefix)
        for seg_text, is_match, is_bracketed in segments:
            run = paragraph.add_run(seg_text)
            if is_match:
                run.bold = True
                run.underline = True
                run.font.highlight_color = WD_COLOR_INDEX.YELLOW
            if is_bracketed:
                run.italic = True
        if reference:
            paragraph.add_run(" ")
            paragraph.add_run(reference)
        paragraph.add_run(" ")
        add_docx_hyperlink(paragraph, "Web", build_reading_link(reference, base_url=base_url))
        return

    if reference:
        paragraph.add_run(reference)
        paragraph.add_run(" ")

    if verse_prefix:
        paragraph.add_run(verse_prefix)

    for seg_text, is_match, is_bracketed in segments:
        run = paragraph.add_run(seg_text)
        if is_match:
            run.bold = True
            run.underline = True
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
        if is_bracketed:
            run.italic = True

    paragraph.add_run(" ")
    add_docx_hyperlink(paragraph, "Web", build_reading_link(reference, base_url=base_url))


def build_xlsx_rich_text(item):
    segments = build_bracketed_segments(item["verse_text"], item.get("highlight_spans", []))
    verse_number = item.get("verse")
    verse_prefix = f"{verse_number}. " if verse_number is not None and verse_number != "" else ""
    rich = CellRichText()

    if verse_prefix:
        rich.append(TextBlock(InlineFont(vertAlign=None), verse_prefix))

    for seg_text, is_match, is_bracketed in segments:
        if not seg_text:
            continue

        font = InlineFont(
            b=bool(is_match),
            i=bool(is_bracketed),
            vertAlign=None,
        )
        rich.append(TextBlock(font, seg_text))

    return rich


def normalize_reading_reference(reference: str) -> str:
    clean_reference = (reference or "").strip()
    if not clean_reference:
        return ""
    clean_reference = clean_reference.replace("[[", "").replace("]]", "").strip()
    if not clean_reference:
        return ""

    chapter_match = re.match(r"^(.+?)\s+(\d+)$", clean_reference, re.IGNORECASE)
    if chapter_match:
        return f"{chapter_match.group(1).strip()} {chapter_match.group(2)}"

    verse_match = re.match(r"^(.+?)\s+(\d+)\s*[:.]\s*(\d+)(?:\s*-\s*\d+)?$", clean_reference, re.IGNORECASE)
    if verse_match:
        return f"{verse_match.group(1).strip()} {verse_match.group(2)}"

    return clean_reference


def build_reading_link(reference: str, base_url: str | None = None) -> str:
    clean_reference = (reference or "").strip()
    if not clean_reference:
        return "#"

    clean_reference = clean_reference.replace("[[", "").replace("]]", "").strip()
    chapter_query = normalize_reading_reference(clean_reference)
    if not chapter_query:
        return "#"

    path = (
        f"/?q={quote_plus(chapter_query)}"
        f"&granularity=Verse&newLineVerse=1&focus={quote_plus(clean_reference)}"
    )
    if base_url:
        base_url = str(base_url).rstrip("/")
        return f"{base_url}{path}"
    return path


def build_markdown_reading_link(reference: str, base_url: str | None = None) -> str:
    clean_reference = (reference or "").strip()
    if not clean_reference:
        return ""

    if clean_reference.startswith("[[") and clean_reference.endswith("]]"):
        return clean_reference
    if clean_reference.startswith("[") and clean_reference.endswith("]"):
        return clean_reference

    return clean_reference


def build_markdown_web_link(reference: str, base_url: str | None = None) -> str:
    clean_reference = (reference or "").strip()
    if not clean_reference:
        return "[Web]()"
    return f"[Web]({build_reading_link(clean_reference, base_url=base_url)})"


@app.get("/", response_class=HTMLResponse)
async def home(request: Request):
    initial_query = request.query_params.get("q", "")
    return templates.TemplateResponse(
        request=request,
        name="index.html",
        context={
            "book_groups": list(BOOK_GROUPS.keys()),
            "initial_query": initial_query,
        }
    )


@app.get("/favicon.ico", include_in_schema=False)
async def favicon():
    return Response(status_code=204)


@app.get("/api/config")
async def config():
    return {
        "book_groups": list(BOOK_GROUPS.keys()),
        "search_types": ["Phrase", "All Words", "Any Word"],
        "reference_styles": ["Wikilink", "Standard"],
        "granularities": ["Verse", "Passage"],
        "reference_positions": ["Ref First", "Ref Last"],
        "context_options": [0, 1, 2, 3, 5, 10]
    }

def build_wildcard_summary(rows, regexes, search_mode):
    summary = {}

    if not regexes:
        return []

    for row in rows:
        verse_text = row["cleaned_text"]
        verse_seen_words = set()

        for rx in regexes:
            for m in rx.finditer(verse_text):
                word = m.group(0)
                if not word:
                    continue

                # Strip leading/trailing punctuation
                clean_word = re.sub(r"^[^\w]+|[^\w]+$", "", word)
                if not clean_word:
                    continue

                key = clean_word.lower()

                if key not in summary:
                    summary[key] = {
                        "word": clean_word,
                        "matches": 0,
                        "verses": 0
                    }

                summary[key]["matches"] += 1

                if key not in verse_seen_words:
                    summary[key]["verses"] += 1
                    verse_seen_words.add(key)

    return sorted(
        summary.values(),
        key=lambda x: (-x["verses"], -x["matches"], x["word"].lower())
    )

@app.post("/api/search")
async def search(payload: SearchRequest):
    try:
        conn = get_conn()
        cur = conn.cursor()

        query = payload.query.strip()
        if not query:
            raise HTTPException(status_code=400, detail="Query is required.")

        if payload.intent == "reference":
            inferred = resolve_reference_intent(query)
            mode = inferred["mode"]
            recommended_view = inferred["recommended_view"]
        else:
            mode = "text"
            recommended_view = "results"
            inferred = {"references": []}

        rows = []
        extracted_refs = []
        regexes = []
        total_matches = 0
        matched_verses = 0
        wildcard_summary = []

        if mode == "text":
            regexes = compile_search_regexes(
                query,
                payload.search_type,
                payload.whole_word,
                payload.case_sensitive
            )

            selected_books = get_selected_book_numbers(payload.book_group)
            use_fts = can_use_fts(cur, query, payload.case_sensitive)

            if use_fts:
                raw_rows = get_fts_text_rows(cur, query, selected_books, payload.search_type)
            else:
                raw_rows = get_sqlite_text_rows(
                    cur, query, selected_books, payload.search_type, payload.case_sensitive
                )

            normalized_rows = [normalize_row(row, payload.suppress_paragraph) for row in raw_rows]
            matched_rows, total_matches = filter_rows_python(normalized_rows, regexes, payload.search_type)
            matched_verses = len(matched_rows)
            if mode == "text":
                wildcard_summary = build_wildcard_summary(matched_rows, regexes, payload.search_type)

            if payload.context_window > 0:
                base_context_rows = [{"id": r["id"]} for r in matched_rows]
                expanded_raw_rows = fetch_context_rows(cur, base_context_rows, payload.context_window)
                rows = [normalize_row(row, payload.suppress_paragraph) for row in expanded_raw_rows]
            else:
                rows = matched_rows

        elif mode == "reference":
            parsed_refs = inferred["references"]
            
            if not parsed_refs:
                raise HTTPException(status_code=400, detail="No valid references found.")

            all_rows = []
            for parsed in parsed_refs:
                all_rows.extend(get_reference_rows(cur, parsed))

            seen = set()
            deduped = []
            for row in all_rows:
                key = (row[COL_ID],)
                if key not in seen:
                    seen.add(key)
                    deduped.append(row)

            expanded_raw_rows = fetch_context_rows(cur, deduped, payload.context_window)
            rows = [normalize_row(row, payload.suppress_paragraph) for row in expanded_raw_rows]

        elif mode == "extract":
            extracted_refs = inferred["references"]

            if payload.lookup_extracted and extracted_refs:
                all_rows = []
                for parsed in extracted_refs:
                    all_rows.extend(get_reference_rows(cur, parsed))

                seen = set()
                deduped = []
                for row in all_rows:
                    key = (row[COL_ID],)
                    if key not in seen:
                        seen.add(key)
                        deduped.append(row)

                expanded_raw_rows = fetch_context_rows(cur, deduped, payload.context_window)
                rows = [normalize_row(row, payload.suppress_paragraph) for row in expanded_raw_rows]
            else:
                rows = []

        else:
            raise HTTPException(status_code=400, detail="Invalid inferred mode.")

        items = build_display_items(
            rows,
            payload.granularity,
            payload.reference_style,
            payload.reference_position,
            regexes,
            mode == "text" and payload.highlight,
            payload.show_red_letters
        )

        reading_items = build_reading_items(
            rows,
            payload.reference_style,
            payload.show_red_letters,
            cur
        )

        conn.close()

        return JSONResponse({
            "query_mode": mode,
            "recommended_view": recommended_view,
            "count": len(items),
            "items": items,
            "reading_items": reading_items,
            "extracted_refs": extracted_refs,
            "wildcard_summary": wildcard_summary,
            "stats": {
                "total_matches": total_matches,
                "matched_verses": matched_verses
            }
        })

    except FileNotFoundError as e:
        raise HTTPException(status_code=500, detail=str(e))
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


def ensure_export_items(items: List[dict] | None):
    if not items:
        raise HTTPException(status_code=400, detail="No results to export.")
    return items


@app.post("/api/export/txt")
async def export_txt(items: List[dict]):
    ensure_export_items(items)
    output = io.StringIO()
    for item in items:
        output.write(build_plain_text(item) + "\n")

    buffer = io.BytesIO(output.getvalue().encode("utf-8"))
    return StreamingResponse(
        buffer,
        media_type="text/plain",
        headers={"Content-Disposition": "attachment; filename=verses.txt"}
    )


@app.post("/api/export/md")
async def export_md(request: Request, items: List[dict]):
    ensure_export_items(items)
    base_url = str(request.base_url).rstrip("/")
    output = io.StringIO()
    for item in items:
        output.write(build_markdown_text(item, base_url=base_url) + "\n")

    buffer = io.BytesIO(output.getvalue().encode("utf-8"))
    return StreamingResponse(
        buffer,
        media_type="text/markdown",
        headers={"Content-Disposition": "attachment; filename=verses.md"}
    )


@app.post("/api/export/csv")
async def export_csv(items: List[dict]):
    ensure_export_items(items)
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(["Reference", "Verse"])
    for item in items:
        writer.writerow([item["reference"], item["verse_text"]])

    buffer = io.BytesIO(output.getvalue().encode("utf-8"))
    return StreamingResponse(
        buffer,
        media_type="text/csv",
        headers={"Content-Disposition": "attachment; filename=verses.csv"}
    )


@app.post("/api/export/docx")
async def export_docx(request: Request, items: List[dict]):
    ensure_export_items(items)
    doc = Document()
    base_url = str(request.base_url).rstrip("/")

    for item in items:
        p = doc.add_paragraph()
        write_docx_runs(p, item, base_url=base_url)

    temp = io.BytesIO()
    doc.save(temp)
    temp.seek(0)

    return StreamingResponse(
        temp,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": "attachment; filename=verses.docx"}
    )


@app.post("/api/export/xlsx")
async def export_xlsx(request: Request, items: List[dict]):
    ensure_export_items(items)
    wb = Workbook()
    ws = wb.active
    ws.title = "Verses"
    base_url = str(request.base_url).rstrip("/")

    ws.append(["Reference", "Web", "Verse"])

    for item in items:
        row_num = ws.max_row + 1
        reference = item.get("reference") or ""

        ws.cell(row=row_num, column=1, value=reference)

        web_cell = ws.cell(row=row_num, column=2, value="Web")
        web_cell.hyperlink = build_reading_link(reference, base_url=base_url)
        web_cell.style = "Hyperlink"

        ws.cell(row=row_num, column=3, value=build_xlsx_rich_text(item))

    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 24
    ws.column_dimensions["C"].width = 120

    temp = io.BytesIO()
    wb.save(temp)
    temp.seek(0)

    return StreamingResponse(
        temp,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=verses.xlsx"}
    )
    
@app.get("/help", response_class=HTMLResponse)
async def help_page(request: Request):
    return templates.TemplateResponse(
        request=request,
        name="help.html",
        context={}
    )
